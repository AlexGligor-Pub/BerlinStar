"""Importul hotelului de anvelope din Excel (Configurări › Import › Hotel anvelope).

Vezi app/services/hotel_import.py. Acoperim: gruparea liniilor in cazari,
normalizarea valorilor (dimensiune, DOT, tip), potrivirea clientului (numar de
masina, apoi nume), crearea cazarii cu anvelope si nomenclatoare, actiunile pe
randurile de rezolvat si revertul.

Rulabil cu pytest sau direct:  python -m tests.test_hotel_import
"""
from __future__ import annotations

import io
from datetime import date

import openpyxl
from sqlalchemy import select

from app.models.anvelopa import Anvelopa
from app.models.cazare_anvelope import CazareAnvelopaItem, CazareAnvelope
from app.models.client import Client
from app.models.client_vehicol import ClientVehicol
from app.models.import_session import ImportRow
from app.models.loc_cazare import LocCazare
from app.models.location import Location
from app.models.marca_anvelopa import MarcaAnvelopa
from app.services import hotel_import as svc
from tests._harness import make_account, make_client, make_receipt, make_session, run

HEADER = ["Nume", "Prenume", "Nr. masina", "Latime", "Inaltime", "Diametru", "Marca", "Profil",
          "Sarcina", "Viteza", "DOT", "Adancime", "Anotimp", "Depozitate", "Depozit"]


def _tire(nume, prenume, plate, data="2025/02/12", loc="C1 DR", **kw):
    return [nume, prenume, plate, kw.get("lat", 215), kw.get("h", 45), kw.get("d", 17), kw.get("marca", "MICHELIN"),
            kw.get("profil", "PRIMACY4"), kw.get("sarcina"), kw.get("viteza"), kw.get("dot", 2722),
            kw.get("adancime", 6), kw.get("anotimp"), data, loc]


def _xlsx(*rows) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Anvelope depozitate"
    ws.append(HEADER)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _stay(nume, prenume, plate, n=4, **kw):
    return [_tire(nume, prenume, plate, **kw) for _ in range(n)]


async def _setup():
    db = await make_session()
    acc = await make_account(db)
    db.add(MarcaAnvelopa(nume="Michelin", status="approved"))
    await db.flush()
    return db, acc


async def _client_with_plate(db, acc, nume, plate):
    c = await make_client(db, acc, nume=nume)
    db.add(ClientVehicol(account_id=acc.id, client_id=c.id, numar_masina=plate))
    await db.flush()
    return c


async def _rows(db, session):
    return (await db.execute(
        select(ImportRow).where(ImportRow.session_id == session.id).order_by(ImportRow.row_number)
    )).scalars().all()


async def _active(db, model, account_id):
    return (await db.execute(
        select(model).where(model.account_id == account_id, model.is_deleted == False).order_by(model.id)
    )).scalars().all()


# ─── Fisierul ─────────────────────────────────────────────────────────────────

async def test_rows_are_grouped_into_stays_and_values_normalized():
    raw = _xlsx(
        *_stay("STEFAN", "ADRIAN", "DJ78AWD"),
        _tire("BADEA", "GABRIEL", "B807EWE", d="16C", dot="2722`", viteza="v", sarcina="X", profil="WINTER CT"),
        _tire("BADEA", "GABRIEL", "B807EWE", lat="X", h="X", d="X", dot=413, adancime="X", anotimp="Vara"),
        _tire(None, None, "VL018215//DJ33DIX", data="nu e data"),
    )
    parsed = svc.parse_file(raw)
    assert parsed.format == "xlsx" and parsed.lines == 7
    assert [len(g.anvelope) for g in parsed.groups] == [4, 2, 1]
    stefan, badea, weird = parsed.groups
    assert (stefan.rows, stefan.data, stefan.depozit) == ([2, 3, 4, 5], date(2025, 2, 12), "C1 DR")
    t1, t2 = badea.anvelope
    assert (t1["dimensiune"], t1["dot"], t1["viteza"], t1["sarcina"], t1["tip"]) == ("215/45 R16C", "2722", "V", None, "iarna")
    assert (t2["dimensiune"], t2["dot"], t2["adancime"], t2["tip"]) == (None, "413", None, "vara")
    assert weird.data is None and svc.plate_keys(weird.numar_masina) == ["VL018215//DJ33DIX", "VL018215", "DJ33DIX"]
    assert any("exact 4" in w for w in parsed.file_warnings)


async def test_file_without_required_columns_is_refused():
    wb = openpyxl.Workbook()
    wb.active.append(["Nume", "Prenume", "Marca"])
    wb.active.append(["A", "B", "C"])
    buf = io.BytesIO()
    wb.save(buf)
    try:
        svc.parse_file(buf.getvalue())
        raise AssertionError("trebuia refuzat")
    except svc.ImportFileError as exc:
        assert "Nr. mașină" in str(exc) and "Depozitate" in str(exc)


# ─── Potrivirea clientului ────────────────────────────────────────────────────

async def test_matching_by_plate_then_name():
    db, acc = await _setup()
    stefan = await _client_with_plate(db, acc, "STEFAN ADRIAN", "DJ78AWD")
    await _client_with_plate(db, acc, "Alt Nume", "DJ11AAA")          # numar gasit, nume diferit
    tancu = await make_client(db, acc, nume="Cristi Tancu")            # doar dupa nume (ordine inversa)
    await make_client(db, acc, nume="BADEA GABRIEL")
    await make_client(db, acc, nume="BADEA GABRIEL")                   # doua persoane cu acelasi nume
    await db.commit()
    rep = await svc.preview(db, acc.id, _xlsx(
        *_stay("STEFAN", "ADRIAN", "DJ78-AWD"),
        *_stay("POPESCU", "ION", "DJ11AAA"),
        *_stay("TANCU", "CRISTI", "DJ33CID"),
        *_stay("BADEA", "GABRIEL", "B807EWE"),
        *_stay("NOU", "CLIENT", "B100NOU"),
        *_stay("STEFAN", "ADRIAN", "DJ78AWD", data="31.02.2025"),
    ))
    assert (rep["cazari"], rep["anvelope"], rep["to_import"]) == (6, 24, 2)
    assert rep["issues_by_type"] == {"error": 1, "duplicate": 0, "new_client": 1, "name_match": 1, "ambiguous": 1}
    by_plate = {i["numar_masina"]: i for i in rep["issues"]}
    assert by_plate["DJ33CID"]["client"] == "Cristi Tancu" and by_plate["DJ33CID"]["candidates"][0]["id"] == tancu.id
    assert len(by_plate["B807EWE"]["candidates"]) == 2
    assert await _active(db, CazareAnvelope, acc.id) == []  # verificarea nu scrie nimic
    assert stefan.id  # folosit doar pentru potrivire


# ─── Importul ─────────────────────────────────────────────────────────────────

async def test_import_creates_stays_tires_and_nomenclatures():
    db, acc = await _setup()
    stefan = await _client_with_plate(db, acc, "STEFAN ADRIAN", "DJ78AWD")
    await make_client(db, acc, nume="TANCU CRISTI")
    await db.commit()
    raw = _xlsx(
        *_stay("STEFAN", "ADRIAN", "DJ78AWD", marca="MICHELIN", profil="PA5", dot=2722),
        _tire("STEFAN", "ADRIAN", "DJ78AWD", data="2024/10/01", loc="MAGAZIE", marca="NECUNOSCUTA SRL"),
        *_stay("TANCU", "CRISTI", "DJ33CID", loc="C3 DR"),
        *_stay("NOU", "CLIENT", "B100NOU", loc="C3 DR"),
    )
    session = await svc.create_session(db, acc.id, raw, filename="custodie.xlsx", actor="manager")
    assert (session.state, session.processed_rows, session.total_rows) == ("done", 4, 4)
    rows = await _rows(db, session)
    assert [(r.status, r.issue) for r in rows] == [
        ("imported", None), ("imported", None), ("pending", "name_match"), ("pending", "new_client"),
    ]

    stays = await _active(db, CazareAnvelope, acc.id)
    assert [(c.client_id, c.numar_masina, c.data_checkin, len(c.items)) for c in stays] == [
        (stefan.id, "DJ78AWD", date(2025, 2, 12), 4), (stefan.id, "DJ78AWD", date(2024, 10, 1), 1),
    ]
    tires = await _active(db, Anvelopa, acc.id)
    assert len(tires) == 5 and all(t.client_id == stefan.id for t in tires)
    assert tires[0].marca_id is not None and tires[0].indice_viteza is None
    # Marca necunoscuta nu se pierde: e propusa spre aprobare si legata de anvelopa.
    propusa = await db.get(MarcaAnvelopa, tires[4].marca_id)
    assert (propusa.nume, propusa.status, propusa.proposed_by_account_id) == ("NECUNOSCUTA SRL", "pending", acc.id)
    assert [l.nume for l in await _active(db, LocCazare, acc.id)] == ["C1 DR", "MAGAZIE"]
    assert rows[0].created["cazare_id"] == stays[0].id and len(rows[0].created["anvelopa_ids"]) == 4
    assert set(session.created) >= {"loc_ids", "dimensiune_ids", "profil_ids", "dot_ids"}


async def test_resolving_pending_stays():
    db, acc = await _setup()
    tancu = await make_client(db, acc, nume="TANCU CRISTI")
    a = await make_client(db, acc, nume="BADEA GABRIEL")
    await make_client(db, acc, nume="BADEA GABRIEL")
    await db.commit()
    session = await svc.create_session(db, acc.id, _xlsx(
        *_stay("TANCU", "CRISTI", "DJ33CID"),
        *_stay("BADEA", "GABRIEL", "B807EWE"),
        *_stay(None, None, "B100NOU", n=2),
    ), filename="f.xlsx", actor="admin")
    confirm, ambiguous, new = await _rows(db, session)
    assert (confirm.issue, ambiguous.issue, new.issue) == ("name_match", "ambiguous", "new_client")

    # Confirmarea potrivirii dupa nume: masina intra si in garajul clientului.
    await svc.import_row(db, confirm, values=None, force_duplicate=False, actor="admin")
    garage = (await db.execute(select(ClientVehicol.numar_masina).where(ClientVehicol.client_id == tancu.id))).scalars().all()
    assert confirm.status == "imported" and garage == ["DJ33CID"]

    # Ambiguu: fara client ales nu se importa; cu clientul ales, da.
    try:
        await svc.import_row(db, ambiguous, values=None, force_duplicate=False, actor="admin")
        raise AssertionError("trebuia sa ceara alegerea clientului")
    except svc.RowActionError:
        pass
    await svc.import_row(db, ambiguous, values={**ambiguous.values, "client_mode": "existing", "client_id": a.id},
                         force_duplicate=False, actor="admin")
    assert ambiguous.status == "imported" and ambiguous.client_id == a.id

    # Client nou fara nume in fisier: primeste un nume-marcaj si masina in garaj.
    result = await svc.import_rows(db, [new], force_duplicates=False, actor="admin")
    assert result == {"imported": 1, "failed": []}
    client = await db.get(Client, new.client_id)
    assert client.nume == "NECOMPLETAT (B100NOU)" and new.created["client_created"] is True

    # Aceeasi masina + aceeasi data a doua oara = duplicat.
    again = await svc.create_session(db, acc.id, _xlsx(*_stay("TANCU", "CRISTI", "DJ33CID")),
                                     filename="g.xlsx", actor="admin", force=True)
    assert (await _rows(db, again))[0].issue == "duplicate"


async def test_brand_and_location_matching_is_tolerant():
    db, acc = await _setup()  # Michelin aprobat
    db.add(MarcaAnvelopa(nume="BFGoodrich", status="approved"))
    db.add(MarcaAnvelopa(nume="Laufenn", status="approved"))
    db.add(MarcaAnvelopa(nume="Propusa", status="pending"))
    db.add(LocCazare(account_id=acc.id, nume="C9"))
    await db.flush()
    n = await svc.Nomenclatoare.load(db, acc.id)
    assert n.marca_id("BF GOODRICH") == n.marca_id("bfgoodrich") is not None
    assert n.marca_id("LAUFEN") is not None
    assert n.marca_id("HANKKOK") is None and n.marca_id("PROPUSA") is None  # fara ghicit, doar aprobate
    assert await n.loc_id("C 9") == await n.loc_id("c9") and n.created["loc_ids"] == []

    # Marcile lipsa se propun; cele deja propuse se refolosesc, cele respinse se repropun.
    noua = await n.ensure_marca("HANKKOK")
    assert (await db.get(MarcaAnvelopa, noua)).status == "pending"
    assert await n.ensure_marca("hankkok") == noua and n.created["marca_ids"] == [noua]
    db.add(MarcaAnvelopa(nume="Respinsa", status="rejected"))
    await db.flush()
    reluata = await n.ensure_marca("RESPINSA")
    marca = await db.get(MarcaAnvelopa, reluata)
    assert (marca.nume, marca.status) == ("Respinsa", "pending")  # numele aprobat de admin ramane


async def test_csv_report_lists_every_row_with_its_outcome():
    db, acc = await _setup()
    await _client_with_plate(db, acc, "STEFAN ADRIAN", "DJ78AWD")
    await db.commit()
    session = await svc.create_session(db, acc.id, _xlsx(
        *_stay("STEFAN", "ADRIAN", "DJ78AWD"),
        *_stay("NOU", "CLIENT", "B100NOU", n=2),
    ), filename="f.xlsx", actor="admin")
    imported, pending = await _rows(db, session)
    lines = [svc.csv_row(r, "STEFAN ADRIAN" if r is imported else None) for r in await _rows(db, session)]
    assert svc.CSV_HEADER[:4] == ["randuri_fisier", "stare", "problema", "nr_masina"]
    assert lines[0][:7] == ["2-5", "imported", "", "DJ78AWD", "STEFAN ADRIAN", "STEFAN ADRIAN", str(imported.client_id)]
    assert lines[0][9:12] == ["4", str(imported.created["cazare_id"]), ""]
    assert lines[1][:6] == ["6-7", "pending", "new_client", "B100NOU", "NOU CLIENT", "nou: NOU CLIENT"]
    assert "client nou" in lines[1][12]
    assert pending.status == "pending"


# ─── Revert ───────────────────────────────────────────────────────────────────

async def test_revert_removes_what_the_import_created_and_keeps_what_changed():
    db, acc = await _setup()
    stefan = await _client_with_plate(db, acc, "STEFAN ADRIAN", "DJ78AWD")
    await db.commit()
    session = await svc.create_session(db, acc.id, _xlsx(
        *_stay("STEFAN", "ADRIAN", "DJ78AWD", loc="C1 DR", profil="PA5"),
        *_stay("STEFAN", "ADRIAN", "DJ78AWD", data="2024/10/01", loc="C2", profil="PA5"),
        *_stay("NOU", "UNU", "B100AAA", loc="C9", profil="DOAR AICI"),
        *_stay("NOU", "DOI", "B200BBB", loc="C9", profil="DOAR AICI"),
    ), filename="f.xlsx", actor="admin")
    r1, r2, new1, new2 = await _rows(db, session)
    await svc.import_rows(db, [new1, new2], force_duplicates=False, actor="admin")

    # Intre timp: a doua cazare a lui Stefan iese din depozit, iar clientul nou 2 primeste un deviz.
    stay2 = await db.get(CazareAnvelope, r2.created["cazare_id"])
    stay2.data_checkout = date(2026, 3, 1)
    receipt = await make_receipt(db, acc)
    receipt.client_id = new2.client_id
    await db.commit()

    plan = await svc.revert_session(db, session, actor="admin", dry_run=True)
    assert (plan["cazari_deleted"], plan["cazari_kept"], plan["anvelope_deleted"]) == (3, 1, 12)
    assert (plan["clients_deleted"], plan["clients_kept"], plan["vehicole_deleted"]) == (1, 1, 1)
    assert len(await _active(db, CazareAnvelope, acc.id)) == 4  # verificarea nu schimba nimic

    await svc.revert_session(db, session, actor="admin", dry_run=False)
    stays = await _active(db, CazareAnvelope, acc.id)
    assert [c.id for c in stays] == [stay2.id]
    assert len(await _active(db, Anvelopa, acc.id)) == 4
    clients = {c.nume for c in await _active(db, Client, acc.id)}
    assert clients == {"STEFAN ADRIAN", "NOU DOI"}
    # Locul C1 DR si C9 nu mai sunt folosite; C2 ramane (cazarea pastrata il foloseste).
    assert [l.nume for l in await _active(db, LocCazare, acc.id)] == ["C2"]
    # Marcile propuse de import si ramase nefolosite dispar odata cu el.
    assert (await db.execute(select(MarcaAnvelopa).where(MarcaAnvelopa.status == "pending",
                                                        MarcaAnvelopa.is_deleted == False))).scalars().all() == []
    await db.refresh(session)
    assert session.state == "reverted"
    assert [r.status for r in await _rows(db, session)] == ["reverted", "imported", "reverted", "reverted"]
    items = (await db.execute(select(CazareAnvelopaItem))).scalars().all()
    assert items  # legaturile raman in tabel; cazarile si anvelopele sunt sterse logic
    assert stefan.id


# ─── Scenarii gasite la code review ───────────────────────────────────────────

async def test_same_person_gets_one_client_when_resolved_in_bulk():
    """Setul de iarna si cel de vara ale aceluiasi om = un singur client nou.

    Fiecare cazare e un rand separat in sesiune; daca fiecare si-ar crea clientul,
    omul ar aparea de doua ori, cu aceeasi masina in doua garaje.
    """
    db, acc = await _setup()
    session = await svc.create_session(db, acc.id, _xlsx(
        *_stay("POPESCU", "ION", "TM01ABC", data="2024/10/01"),
        *_stay("POPESCU", "ION", "TM01ABC", data="2025/04/01"),
        *_stay("POPESCU", "ION", "TM02XYZ", data="2025/04/01"),   # acelasi om, alta masina
    ), filename="f.xlsx", actor="admin")
    rows = await _rows(db, session)
    assert [r.issue for r in rows] == ["new_client"] * 3

    assert (await svc.import_rows(db, rows, force_duplicates=False, actor="admin"))["imported"] == 3
    clients = await _active(db, Client, acc.id)
    assert [c.nume for c in clients] == ["POPESCU ION"]
    assert len({r.client_id for r in await _rows(db, session)}) == 1
    garage = (await db.execute(
        select(ClientVehicol.numar_masina).where(ClientVehicol.client_id == clients[0].id)
    )).scalars().all()
    assert sorted(garage) == ["TM01ABC", "TM02XYZ"]
    assert len(await _active(db, CazareAnvelope, acc.id)) == 3


async def test_name_only_on_the_first_line_keeps_the_set_together():
    """In fisierele reale numele e scris o singura data, pe prima linie a setului."""
    parsed = svc.parse_file(_xlsx(
        _tire("POPESCU", "ION", "TM01ABC"),
        _tire(None, None, "TM01ABC"),
        _tire(None, None, "TM01ABC"),
        _tire(None, None, "TM01ABC"),
    ))
    assert len(parsed.groups) == 1
    g = parsed.groups[0]
    assert (len(g.anvelope), g.full_name) == (4, "POPESCU ION")


async def test_different_names_on_the_same_stay_are_reported():
    parsed = svc.parse_file(_xlsx(
        _tire("POPESCU", "ION", "TM01ABC"),
        _tire("IONESCU", "DAN", "TM01ABC"),
    ))
    g = parsed.groups[0]
    assert len(parsed.groups) == 1 and g.full_name == "POPESCU ION"
    assert any("nume diferit" in m for m in g.messages)


async def test_hidden_sheets_are_ignored():
    """Registrele exportate din aplicatii vechi au foi ascunse cu arhive."""
    wb = openpyxl.Workbook()
    arhiva = wb.active
    arhiva.title = "Arhiva"
    arhiva.append(HEADER)
    arhiva.append(_tire("VECHI", "CLIENT", "XX11XXX"))
    arhiva.sheet_state = "hidden"
    curent = wb.create_sheet("Curent")
    curent.append(HEADER)
    for r in _stay("NOU", "CLIENT", "TM01ABC"):
        curent.append(r)
    buf = io.BytesIO()
    wb.save(buf)

    parsed = svc.parse_file(buf.getvalue())
    assert parsed.sheet == "Curent"
    assert [g.numar_masina for g in parsed.groups] == ["TM01ABC"]


async def test_messy_plate_is_cleaned_before_it_reaches_the_garage():
    """„VL018215//DJ33DIX" gaseste clientul, dar in garaj si pe cazare scriem
    numarul real — altfel POS-ul nu mai gaseste masina dupa placuta."""
    db, acc = await _setup()
    client = await _client_with_plate(db, acc, "ION POPA", "DJ33DIX")
    await db.commit()
    session = await svc.create_session(
        db, acc.id, _xlsx(*_stay("ION", "POPA", "VL018215//DJ33DIX")), filename="f.xlsx", actor="admin",
    )
    row = (await _rows(db, session))[0]
    assert (row.status, row.client_id) == ("imported", client.id)
    garage = (await db.execute(
        select(ClientVehicol.numar_masina).where(
            ClientVehicol.client_id == client.id, ClientVehicol.is_deleted == False)
    )).scalars().all()
    assert garage == ["DJ33DIX"]  # nu s-a mai adaugat o masina
    stay = (await _active(db, CazareAnvelope, acc.id))[0]
    assert stay.numar_masina == "DJ33DIX" and "VL018215//DJ33DIX" in stay.comments


async def test_revert_keeps_a_stay_whose_tire_was_edited():
    db, acc = await _setup()
    await _client_with_plate(db, acc, "STEFAN ADRIAN", "DJ78AWD")
    await db.commit()
    session = await svc.create_session(
        db, acc.id, _xlsx(*_stay("STEFAN", "ADRIAN", "DJ78AWD")), filename="f.xlsx", actor="admin",
    )
    row = (await _rows(db, session))[0]
    tire = await db.get(Anvelopa, row.created["anvelopa_ids"][0])
    tire.adancime, tire.updated_at = 3.5, svc.common.now()   # mecanicul a masurat-o
    await db.commit()

    plan = await svc.revert_session(db, session, actor="admin", dry_run=True)
    assert (plan["cazari_deleted"], plan["cazari_kept"]) == (0, 1)
    assert "anvelopă" in plan["kept"][0]["reason"]
    await svc.revert_session(db, session, actor="admin", dry_run=False)
    assert len(await _active(db, CazareAnvelope, acc.id)) == 1
    assert len(await _active(db, Anvelopa, acc.id)) == 4


async def test_revert_returns_a_reproposed_brand_to_rejected():
    """Marca respinsa exista GLOBAL dinaintea importului: la revert se intoarce la
    `rejected`. Stearsa, ar bloca numele pentru toate conturile."""
    db, acc = await _setup()
    db.add(MarcaAnvelopa(nume="Sailun", status="rejected"))
    await _client_with_plate(db, acc, "STEFAN ADRIAN", "DJ78AWD")
    await db.commit()
    session = await svc.create_session(
        db, acc.id, _xlsx(*_stay("STEFAN", "ADRIAN", "DJ78AWD", marca="SAILUN")), filename="f.xlsx", actor="admin",
    )
    marca = (await db.execute(select(MarcaAnvelopa).where(MarcaAnvelopa.nume == "Sailun"))).scalar_one()
    assert marca.status == "pending"

    await svc.revert_session(db, session, actor="admin", dry_run=False)
    await db.refresh(marca)
    assert (marca.status, marca.is_deleted) == ("rejected", False)


# ─── Punctul de lucru ─────────────────────────────────────────────────────────

async def _location(db, acc, name):
    loc = Location(account_id=acc.id, name=name)
    db.add(loc)
    await db.flush()
    return loc


async def test_stays_land_on_the_single_location_of_the_account():
    """Pagina Hotel arata doar cazarile locatiei statiei curente, deci o cazare
    fara `location_id` n-ar fi vizibila nicaieri. Cu un singur punct de lucru nu
    e nimic de ales: importul il pune singur — si la procesare, si la randurile
    rezolvate mai tarziu."""
    db, acc = await _setup()
    loc = await _location(db, acc, "Craiova")
    await _client_with_plate(db, acc, "STEFAN ADRIAN", "DJ78AWD")
    await db.commit()
    session = await svc.create_session(
        db, acc.id,
        _xlsx(*_stay("STEFAN", "ADRIAN", "DJ78AWD"), *_stay("NOU", "CLIENT", "B100NOU")),
        filename="f.xlsx", actor="admin",
    )
    assert session.location_id == loc.id
    nou = [r for r in await _rows(db, session) if r.issue == "new_client"][0]
    await svc.import_row(db, nou, values=None, force_duplicate=False, actor="admin")

    cazari = await _active(db, CazareAnvelope, acc.id)
    assert len(cazari) == 2
    assert [c.location_id for c in cazari] == [loc.id, loc.id]


async def test_several_locations_require_an_explicit_choice():
    """Cu mai multe puncte de lucru alegerea e a utilizatorului: fara ea importul
    se opreste inainte sa scrie ceva, iar o locatie din alt cont e refuzata."""
    db, acc = await _setup()
    await _location(db, acc, "Craiova")
    bucuresti = await _location(db, acc, "București")
    alt_cont = await make_account(db, "alta", "alta")
    strain = await _location(db, alt_cont, "A altora")
    await db.commit()
    raw = _xlsx(*_stay("NOU", "CLIENT", "B100NOU"))

    for location_id, asteptat in ((None, "mai multe puncte"), (strain.id, "nu există")):
        try:
            await svc.create_session(db, acc.id, raw, filename="f.xlsx", actor="admin", location_id=location_id)
            raise AssertionError("trebuia refuzat")
        except svc.ImportFileError as exc:
            assert asteptat in str(exc), str(exc)
    assert await _active(db, CazareAnvelope, acc.id) == []

    session = await svc.create_session(
        db, acc.id, raw, filename="f.xlsx", actor="admin", location_id=bucuresti.id,
    )
    row = (await _rows(db, session))[0]
    await svc.import_row(db, row, values=None, force_duplicate=False, actor="admin")
    assert [c.location_id for c in await _active(db, CazareAnvelope, acc.id)] == [bucuresti.id]


TESTS = [v for k, v in sorted(globals().items()) if k.startswith("test_")]

if __name__ == "__main__":
    for t in TESTS:
        run(t())
    print(f"OK — {len(TESTS)} scenarii de import hotel anvelope trecute.")
